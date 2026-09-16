"""Tests for branded-deck revision helpers (TDD)."""

from __future__ import annotations

import sys
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Inches

REPO_ROOT = Path(__file__).resolve().parent.parent
SCRIPTS = REPO_ROOT / "scripts"
sys.path.insert(0, str(SCRIPTS))

from revise_wings3_branded_deck import (  # noqa: E402
    BRANDED_PPTX,
    LADDER_LINES,
    add_screenshot_placeholder,
    apply_text_replacements,
    insert_content_slide,
    move_slide,
    slide_text,
)

pytest.importorskip("pptx")


def _blank_deck() -> Presentation:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Without a tracking server on the cluster you cannot operate the agent"
    body = slide.placeholders[1]
    tf = body.text_frame
    tf.text = "Without traces you cannot fix it"
    return prs


def test_ladder_lines_are_constructive_and_numbered():
    assert len(LADDER_LINES) == 4
    joined = " ".join(LADDER_LINES).lower()
    assert "without a tracking server" not in joined
    assert "you can see" in joined
    assert "you can fix" in joined
    assert "you can prove" in joined
    assert "you can ship" in joined
    for i, line in enumerate(LADDER_LINES, 1):
        assert line.startswith(f"{i}.")


def test_apply_text_replacements_rewrites_ladder():
    prs = _blank_deck()
    apply_text_replacements(prs)
    blob = slide_text(prs.slides[0]).lower()
    assert "without a tracking server on the cluster" not in blob
    assert "you can see the agent" in blob
    assert "you can fix it" in blob


def test_screenshot_placeholder_contains_capture_and_why():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    add_screenshot_placeholder(
        slide,
        what="OpenShift AI → Projects → my-first-model",
        why="this namespace is the MLflow workspace",
    )
    blob = slide_text(slide)
    assert "<screenshot>" in blob
    assert "What to capture: OpenShift AI → Projects → my-first-model" in blob
    assert "Why it is here: this namespace is the MLflow workspace" in blob


def test_insert_and_move_slide_puts_new_slide_after_title():
    prs = Presentation()
    title = prs.slides.add_slide(prs.slide_layouts[0])
    title.shapes.title.text = "Title"
    second = prs.slides.add_slide(prs.slide_layouts[1])
    second.shapes.title.text = "Old second"
    insert_content_slide(prs, "Terms", ["Project — namespace"], notes="Define first.")
    move_slide(prs, len(prs.slides) - 1, 1)
    assert prs.slides[1].shapes.title.text == "Terms"
    assert "Project — namespace" in slide_text(prs.slides[1])
    assert prs.slides[2].shapes.title.text == "Old second"


@pytest.mark.skipif(not BRANDED_PPTX.is_file(), reason="branded pptx is gitignored")
def test_branded_deck_exists_for_live_revision():
    assert BRANDED_PPTX.stat().st_size > 1_000_000


@pytest.mark.skipif(not BRANDED_PPTX.is_file(), reason="branded pptx is gitignored")
def test_revise_branded_deck_applies_feedback(tmp_path):
    import shutil

    from revise_wings3_branded_deck import revise

    dest = tmp_path / "deck.pptx"
    shutil.copy2(BRANDED_PPTX, dest)
    prs = Presentation(str(dest))
    revise(prs)
    blob = "\n".join(slide_text(s) for s in prs.slides)
    assert "Six words you will reuse all hour" in blob
    assert "OpenShift AI — Projects" in blob
    assert "Launch MLflow — recap what's what" in blob
    assert "<screenshot>" in blob
    assert "What to capture:" in blob
    assert "Agent observability with MLflow on OpenShift AI" in blob
    assert "you can see the agent" in blob
    assert "Data Scientist" not in blob
    assert "three hats" not in blob.lower()
    assert "opendatahub.io/mlflow-instance" in blob
    assert any("Judge" in slide_text(s) for s in prs.slides)
    first_lines = [slide_text(s).split("\n")[0][:80] for s in prs.slides]
    joined = "\n".join(first_lines[:15])
    assert "Six words you will reuse all hour" in joined
