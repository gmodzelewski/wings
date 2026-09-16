"""Revise the branded WINGS3 Google-exported deck from presentation feedback.

Edits AI Wings 3 - Deep Dive.pptx in place (one .bak copy).
New UI shots are <screenshot> placeholders — replace them when a cluster is up.

Run: python3 scripts/revise_wings3_branded_deck.py
"""

from __future__ import annotations

import shutil
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.util import Inches, Pt

_SCRIPTS_DIR = Path(__file__).resolve().parent
WINGS3_ROOT = _SCRIPTS_DIR.parent
BRANDED_PPTX = WINGS3_ROOT / "AI Wings 3 - Deep Dive.pptx"

LADDER_LINES = [
    "1. Tracking server on the platform — you can see the agent",
    "2. Traces — you can fix it",
    "3. Evaluation — you can prove a prompt change helped",
    "4. Dataset + judge — you can ship with a gate you can argue with",
]

LADDER_OLD_TO_NEW = [
    (
        "Without a tracking server on the cluster you cannot operate the agent",
        LADDER_LINES[0],
    ),
    ("Without traces you cannot fix it", LADDER_LINES[1]),
    (
        "Without eval you cannot prove a prompt change helped",
        LADDER_LINES[2],
    ),
    (
        "Without a golden set and a judge you can argue with, you cannot ship",
        LADDER_LINES[3],
    ),
    (
        "Without a golden dataset and a judge you can argue with, you cannot ship",
        LADDER_LINES[3],
    ),
]

TERMS_BULLETS = [
    "Project — OpenShift namespace my-first-model (dashboard)",
    "Workspace — MLflow name for that project (MLFLOW_WORKSPACE). The RBAC boundary.",
    "Experiment — named bucket inside the workspace",
    "Trace — one recorded agent run: LLM calls, tool spans, timeline",
    "Dataset — named, versioned inputs plus expected answers you re-run (golden set math_golden)",
    "Judge — a scorer that is itself an LLM, with a written rationale, not a substring check",
]

TWO_ROLES_BODY = (
    "Platform engineer — installed MLflow (operator + CR)\n"
    "AI engineer — workbench authorized to the workspace\n"
    "Acts 2–4 are the same AI engineer\n"
    "Q&A"
)

WHY_NATIVE_BULLETS = [
    "autolog works against local MLflow or a SaaS tracer — no OpenShift instance required",
    "Challenges: extra identity, a laptop token, no project RBAC, traces live somewhere else",
    "Native MLflow: workspace is the project, URI injected, Kubernetes auth",
    "Act 1 makes this concrete in the OpenShift AI and MLflow UIs",
]


def _iter_shapes(shapes):
    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            try:
                yield from _iter_shapes(shape.shapes)
            except Exception:
                continue


def slide_text(slide) -> str:
    parts = []
    for shape in _iter_shapes(slide.shapes):
        if shape.has_text_frame:
            t = shape.text_frame.text.strip()
            if t:
                parts.append(t)
    return "\n".join(parts)


def _set_frame_text(tf, text: str) -> None:
    paras = list(tf.paragraphs)
    lines = text.split("\n")
    if not paras:
        tf.text = text
        return
    for i, line in enumerate(lines):
        if i < len(paras):
            p = paras[i]
        else:
            p = tf.add_paragraph()
        if p.runs:
            p.runs[0].text = line
            for run in p.runs[1:]:
                run.text = ""
        else:
            p.text = line
    for p in paras[len(lines) :]:
        if p.runs:
            for run in p.runs:
                run.text = ""
        else:
            p.text = ""


def _replace_in_frame(tf, old: str, new: str) -> bool:
    full = tf.text
    if old not in full:
        return False
    _set_frame_text(tf, full.replace(old, new))
    return True


def apply_text_replacements(prs: Presentation) -> None:
    for slide in prs.slides:
        for shape in _iter_shapes(slide.shapes):
            if not shape.has_text_frame:
                continue
            t = shape.text_frame.text
            if t.strip() == "MLFlow":
                _set_frame_text(
                    shape.text_frame,
                    "Agent observability with MLflow on OpenShift AI",
                )
                continue
            if "four acts and three hats" in t:
                _set_frame_text(
                    shape.text_frame,
                    "Two roles — then we stop talking about hats",
                )
                continue
            if (
                "Act 1 install verify - Platform engineer" in t
                and "Data Scientist" in t
                and "Q&A" in t
            ):
                _set_frame_text(shape.text_frame, TWO_ROLES_BODY)
                continue
            if "AI Developer" in t and "Act 2" in t and "Q&A" not in t and "Act 1" not in t:
                _set_frame_text(shape.text_frame, "Act 2 autolog traces")
                continue
            if "Data Scientist" in t and "Act 3" in t and "Q&A" not in t:
                _set_frame_text(shape.text_frame, "Act 3 evaluation")
                continue
            if "Data Scientist" in t and "Act 4" in t and "Q&A" not in t:
                _set_frame_text(shape.text_frame, "Act 4 production-grade evals")
                continue
            if "Installation is supereasy" in t:
                _set_frame_text(
                    shape.text_frame,
                    "Operator-managed instance\n"
                    "Tracking server on the platform — you can see the agent",
                )
                continue
            if "We are using OpenShift AI" in t:
                _set_frame_text(
                    shape.text_frame,
                    "Tracing works without a cluster instance. "
                    "OpenShift AI includes a managed tracking server — that is the platform story.",
                )
                continue
            if "MLFlow envs injected automatically" in t:
                _set_frame_text(
                    shape.text_frame,
                    "UI workbenches: annotation is automatic after install",
                )
                continue
            if "opendatahub.io/mlflow-instance=mlflow on the Notebook injects" in t:
                _set_frame_text(
                    shape.text_frame,
                    "GitOps / YAML must set opendatahub.io/mlflow-instance=mlflow. "
                    "UI workbenches created after install get it automatically. It injects:\n"
                    "- MLFLOW_TRACKING_URI\n"
                    "- MLFLOW_K8S_INTEGRATION=true\n"
                    "- MLFLOW_TRACKING_AUTH=kubernetes-namespaced",
                )
                continue
            for old, new in LADDER_OLD_TO_NEW:
                if old in t:
                    _replace_in_frame(shape.text_frame, old, new)
                    t = shape.text_frame.text
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame
            nt = notes.text
            for old, new in LADDER_OLD_TO_NEW:
                if old in nt:
                    _replace_in_frame(notes, old, new)
                    nt = notes.text


def _layout(prs: Presentation, *names):
    by_name = {layout.name: layout for layout in prs.slide_layouts}
    for name in names:
        if name in by_name:
            return by_name[name]
    return prs.slide_layouts[0]


def add_screenshot_placeholder(
    slide,
    what: str,
    why: str,
    left: float = 0.6,
    top: float = 3.5,
    width: float = 12.1,
    height: float = 3.5,
) -> None:
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(0xF5, 0xF5, 0xF5)
    shape.line.color.rgb = RGBColor(0xCC, 0x00, 0x00)
    tf = shape.text_frame
    tf.word_wrap = True
    tf.text = f"<screenshot>\nWhat to capture: {what}\nWhy it is here: {why}"
    for para in tf.paragraphs:
        for run in para.runs:
            run.font.size = Pt(16)
            run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)


def insert_content_slide(
    prs: Presentation,
    title: str,
    bullets: list[str],
    notes: str = "",
) -> object:
    layout = _layout(prs, "TITLE_ONLY", "BLANK", "TITLE")
    slide = prs.slides.add_slide(layout)
    if slide.shapes.title is not None:
        slide.shapes.title.text = title
    else:
        box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12.2), Inches(0.8))
        p = box.text_frame.paragraphs[0]
        p.text = title
        p.font.size = Pt(28)
        p.font.bold = True
    body = slide.shapes.add_textbox(Inches(0.5), Inches(1.2), Inches(12.2), Inches(2.2))
    tf = body.text_frame
    tf.word_wrap = True
    for i, line in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.level = 0
        p.font.size = Pt(18)
    if notes:
        slide.notes_slide.notes_text_frame.text = notes
    return slide


def move_slide(prs: Presentation, old_index: int, new_index: int) -> None:
    xml_slides = prs.slides._sldIdLst  # noqa: SLF001
    slides = list(xml_slides)
    el = slides[old_index]
    xml_slides.remove(el)
    xml_slides.insert(new_index, el)


def _index_containing(prs: Presentation, needle: str, exclude: str | None = None) -> int:
    for i, slide in enumerate(prs.slides):
        text = slide_text(slide)
        if needle in text and (exclude is None or exclude not in text):
            return i
    raise KeyError(needle)


def _already_revised(prs: Presentation) -> bool:
    return any("Six words you will reuse all hour" in slide_text(s) for s in prs.slides)


def _append_and_move(prs: Presentation, after_index: int) -> None:
    move_slide(prs, len(prs.slides) - 1, after_index + 1)


def _place_embedded_experiments(prs: Presentation) -> None:
    """Keep the existing OpenShift AI Experiments screenshot in the product tour."""
    recap = _index_containing(prs, "Launch MLflow — recap")
    try:
        create = _index_containing(prs, "Create instance:")
        docs = _index_containing(prs, "Further docs:")
    except KeyError:
        return
    pic = None
    for i, slide in enumerate(prs.slides):
        if create < i < docs and not slide_text(slide).strip():
            pic = i
            break
    if pic is None:
        return
    if pic == recap:
        return
    move_slide(prs, pic, recap)


def revise(prs: Presentation) -> Presentation:
    if _already_revised(prs):
        apply_text_replacements(prs)
        _place_embedded_experiments(prs)
        return prs

    apply_text_replacements(prs)

    insert_content_slide(
        prs,
        "Six words you will reuse all hour",
        TERMS_BULLETS,
        notes=(
            "Say these once, before the ladder uses dataset and judge. "
            "Project is the namespace. Workspace is MLFLOW_WORKSPACE=my-first-model. "
            "A trace is one run. A dataset is the yardstick. A judge writes why a row passed or failed."
        ),
    )
    _append_and_move(prs, _index_containing(prs, "AI Wings 3"))

    insert_content_slide(
        prs,
        "Tracing works without a cluster instance — here is why we still install one",
        WHY_NATIVE_BULLETS,
        notes=(
            "Pick this up before any YAML. Tracing is not invented by the operator. "
            "Contrast an external tracer: extra identity, extra URL, a token on the laptop. "
            "Here the workspace is the project, the URI is injected, and you stay inside OpenShift AI."
        ),
    )
    _append_and_move(
        prs,
        _index_containing(prs, "Tracing works without a cluster instance. OpenShift AI"),
    )

    act1 = _index_containing(prs, "Act 1 install verify", exclude="Act 2")
    insert_content_slide(
        prs,
        "OpenShift AI — Projects",
        ["Start in the dashboard. Project my-first-model is the MLflow workspace."],
        notes="Do not start at oc patch. Step-by-step for less prior knowledge.",
    )
    add_screenshot_placeholder(
        prs.slides[-1],
        what="OpenShift AI → Projects → my-first-model",
        why="this namespace is the MLflow workspace",
    )
    _append_and_move(prs, act1)

    insert_content_slide(
        prs,
        "OpenShift AI — Project page",
        [
            "Dashboard-created workbenches: annotation is automatic after install",
            "GitOps / YAML Notebook: set opendatahub.io/mlflow-instance=mlflow yourself",
        ],
        notes=(
            "This hour's workbench is GitOps (workbench-wings3-demo.yaml) so the annotation "
            "is in the manifest. If they Create workbench from the UI after MLflow exists, "
            "the platform sets it for them."
        ),
    )
    add_screenshot_placeholder(
        prs.slides[-1],
        what="OpenShift AI → Project my-first-model → Workbenches (wings3-demo Running)",
        why=(
            "UI workbenches created after MLflow install get "
            "opendatahub.io/mlflow-instance automatically; GitOps YAML must set it"
        ),
    )
    _append_and_move(prs, act1 + 1)

    insert_content_slide(
        prs,
        "Launch MLflow — recap what's what",
        [
            "Experiments — run buckets (classic tracking; not the deep dive)",
            "Traces — agent observability (Act 2)",
            "Evaluation — scored runs (Act 3)",
            "Datasets — named golden sets (Act 4)",
        ],
        notes=(
            "Keep the embedded Experiments (MLflow) screenshot that follows. "
            "Then Launch MLflow. Recap the four surfaces so experiment tracking is named, "
            "not skipped, and the hour still goes deep on traces and eval."
        ),
    )
    add_screenshot_placeholder(
        prs.slides[-1],
        what="Launch MLflow → standalone /mlflow home, workspace dropdown my-first-model",
        why="this is the UI for Traces, Evaluation, and Datasets — not the embedded Experiments list",
    )
    _append_and_move(prs, act1 + 2)

    _place_embedded_experiments(prs)
    return prs


def main() -> None:
    if not BRANDED_PPTX.is_file():
        print(f"missing {BRANDED_PPTX}", file=sys.stderr)
        sys.exit(1)
    bak = BRANDED_PPTX.with_suffix(".pptx.bak")
    if not bak.exists():
        shutil.copy2(BRANDED_PPTX, bak)
        print(f"backup {bak}")
    prs = Presentation(str(BRANDED_PPTX))
    revise(prs)
    prs.save(str(BRANDED_PPTX))
    print(f"wrote {BRANDED_PPTX} ({len(prs.slides)} slides)")


if __name__ == "__main__":
    main()
